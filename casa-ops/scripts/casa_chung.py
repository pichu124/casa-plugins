"""Hàm dùng chung cho các script của plugin casa-ops: chuẩn hóa chữ, giờ, SĐT, cơ sở, tuổi bé,
đọc CSV/XLSX và trình bày bảng. Quy tắc chuẩn hóa mô tả ở references/du-lieu-lead.md.

Chỉ dùng thư viện chuẩn của Python. Mọi giờ là giờ Việt Nam, không múi giờ.
"""

from __future__ import annotations

import csv
import re
import sys
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path


CO_SO = ["Chùa Láng", "Mỹ Đình", "Kim Mã", "Quốc tế – chưa rõ cơ sở", "Chưa rõ"]


NHOM_TUOI = [
    "15 tháng – <2 tuổi",
    "2 – <3 tuổi",
    "3 – <5 tuổi",
    "5 – 6 tuổi",
    "Ngoài độ tuổi",
    "Chưa rõ",
]


def bo_dau(s: str) -> str:
    """Chữ thường, bỏ dấu tiếng Việt, gộp khoảng trắng: để so khớp mềm."""
    s = unicodedata.normalize("NFD", s or "")
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "D").lower()
    return re.sub(r"\s+", " ", s).strip()


def doc_gio(s: str) -> datetime | None:
    s = (s or "").strip()
    if not s:
        return None
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%dT%H:%M:%S",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%Y-%m-%d",
        "%d/%m/%Y",
    ):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def chu_so(s: str) -> str:
    return re.sub(r"\D", "", s or "")


def co_sdt(s: str) -> bool:
    return len(chu_so(s)) >= 9


def che_sdt(s: str) -> str:
    so = chu_so(s)
    return f"…{so[-4:]}" if len(so) >= 4 else "…"


def ten_hien_thi(ten: str) -> str:
    ten = re.sub(r"\s+", " ", (ten or "")).strip()
    if not ten or bo_dau(ten) == "khach chua ro ten":
        return "Khách chưa rõ tên"
    return ten


def chuan_hoa_co_so(s: str) -> str:
    t = bo_dau(s)
    if not t:
        return "Chưa rõ"
    nhom = {
        "Chùa Láng": ["chua lang", "hongkong", "hong kong", "huynh thuc khang", "hk"],
        "Mỹ Đình": ["my dinh", "flc", "pham hung"],
        "Kim Mã": ["kim ma", "van phuc", "song ngu", "ngoc ha"],
    }
    trung = [
        ten
        for ten, tu in nhom.items()
        if any(re.search(rf"(?<![a-z]){re.escape(x)}(?![a-z])", t) for x in tu)
    ]
    if len(trung) == 1:
        return trung[0]
    if len(trung) > 1:
        return "Chưa rõ"  # Phụ huynh nhắc nhiều cơ sở: không đoán.
    if re.search(r"quoc te|international", t):
        return "Quốc tế – chưa rõ cơ sở"
    return "Chưa rõ"


SO_CHU = {"mot": 1, "hai": 2, "ba": 3, "bon": 4, "tu": 4, "nam": 5, "sau": 6, "bay": 7}


def _so(x: str) -> float:
    x = x.replace(",", ".")
    return float(SO_CHU[x]) if x in SO_CHU else float(x)


def khoang_thang(tuoi: str, ngay_tinh: datetime) -> tuple[float, float] | None:
    """Khoảng tuổi (tháng) [thấp, cao) mà chữ phụ huynh viết có thể ứng với. None nếu không hiểu."""
    t = bo_dau(tuoi)
    if not t:
        return None
    so = r"(\d+(?:[.,]\d+)?|mot|hai|ba|bon|tu|nam|sau|bay)"

    # Năm sinh: "sinh năm 2022", "sn 2022", "2022"; có tháng: "sinh tháng 3/2023", "03/2023".
    m = re.search(r"(?:thang\s*)?(\d{1,2})\s*[/-]\s*(20\d{2})", t)
    if m:
        thang, nam = int(m.group(1)), int(m.group(2))
        if 1 <= thang <= 12:
            tuoi_thang = (ngay_tinh.year - nam) * 12 + (ngay_tinh.month - thang)
            return (tuoi_thang, tuoi_thang + 1)
    m = re.search(r"(?<!\d)(20\d{2})(?!\d)", t)
    if m:
        nam = int(m.group(1))
        cao = (ngay_tinh.year - nam) * 12 + ngay_tinh.month  # sinh đầu năm
        thap = cao - 12  # sinh cuối năm
        return (max(thap, 0), cao)

    duoi = bool(re.search(r"\b(duoi|chua (du|toi|den))\b", t))

    # "3-4 tuổi", "3 đến 4 tuổi"
    m = re.search(rf"{so}\s*(?:-|–|den|toi)\s*{so}\s*tuoi", t)
    if m:
        a, b = _so(m.group(1)), _so(m.group(2))
        return (a * 12, (b + 1) * 12)

    # "x tuổi y tháng", "x tuổi rưỡi", "x tuổi"; "2t" là 2 tuổi ("18th" là tháng, xử lý dưới).
    t = re.sub(r"(\d)\s*t\b", r"\1 tuoi", t)
    m = re.search(rf"{so}\s*tuoi(?:\s*(ruoi)|\s*{so}\s*thang)?", t)
    if m:
        nam = _so(m.group(1))
        if m.group(2):
            return (nam * 12 + 6, nam * 12 + 7)
        if m.group(3):
            thang = _so(m.group(3))
            return (nam * 12 + thang, nam * 12 + thang + 1)
        if duoi:
            return (max(nam * 12 - 12, 0), nam * 12)
        if nam != int(nam):  # "1,5 tuổi"
            return (nam * 12, nam * 12 + 1)
        return (nam * 12, nam * 12 + 12)

    # "18 tháng", "dưới 18 tháng"
    m = re.search(rf"{so}\s*(?:thang|th\b)", t)
    if m:
        thang = _so(m.group(1))
        if duoi:
            return (max(thang - 6, 0), thang)
        return (thang, thang + 1)
    return None


def nhom_tuoi(tuoi: str, ngay_tinh: datetime) -> str:
    khoang = khoang_thang(tuoi, ngay_tinh)
    if khoang is None:
        return "Chưa rõ"
    giua = (khoang[0] + khoang[1]) / 2  # Lấy điểm giữa của khoảng có thể.
    # Trường nhận bé tới 6 tuổi: "6 tuổi" (72–83 tháng) vẫn thuộc nhóm 5 – 6 tuổi.
    if giua < 15 or giua >= 84:
        return "Ngoài độ tuổi"
    if giua < 24:
        return "15 tháng – <2 tuổi"
    if giua < 36:
        return "2 – <3 tuổi"
    if giua < 60:
        return "3 – <5 tuổi"
    return "5 – 6 tuổi"


def doc_hang(duong_dan: Path, ten_tab: str = "Leads") -> list[dict[str, str]]:
    """Đọc CSV hoặc XLSX (tab `ten_tab`, không có thì tab đầu) thành list dict theo hàng tiêu đề."""
    if duong_dan.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl  # type: ignore
        except ImportError:
            sys.exit(f"Cần openpyxl để đọc .xlsx (pip install openpyxl), hoặc xuất tab {ten_tab} ra .csv.")
        ws = openpyxl.load_workbook(duong_dan, read_only=True, data_only=True)
        sheet = ws[ten_tab] if ten_tab in ws.sheetnames else ws.worksheets[0]
        hang = list(sheet.iter_rows(values_only=True))
        if not hang:
            return []
        tieu_de = ["" if c is None else str(c) for c in hang[0]]

        def o(c):
            if c is None:
                return ""
            if isinstance(c, datetime):
                return c.strftime("%Y-%m-%d %H:%M:%S")
            return str(c)

        return [dict(zip(tieu_de, (o(c) for c in r))) for r in hang[1:]]
    with open(duong_dan, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def pt(x: float | None) -> str:
    if x is None:
        return "–"
    s = f"{x * 100:.1f}".replace(".", ",")
    return (s[:-2] if s.endswith(",0") else s) + "%"


def bang(tieu_de: list[str], hang: list[list]) -> str:
    dong = ["| " + " | ".join(tieu_de) + " |", "|" + "|".join("---" for _ in tieu_de) + "|"]
    dong += ["| " + " | ".join(str(o) for o in h) + " |" for h in hang]
    return "\n".join(dong)


def gio_viet_nam() -> datetime:
    try:
        from zoneinfo import ZoneInfo

        return datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).replace(tzinfo=None)
    except Exception:
        return datetime.utcnow() + timedelta(hours=7)


SDT_TRONG_CHU = re.compile(r"(?<!\d)(?:\+?84|0)(?:[ .\-]?\d){8,10}(?!\d)")
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")


def an_thong_tin(s: str, toi_da: int = 200) -> str:
    """Thay SĐT bằng [SĐT], email bằng [email], gộp khoảng trắng, cắt bớt cho gọn.
    Tên bé, địa chỉ cụ thể không lọc được bằng máy: người trình bày vẫn phải đọc lại."""
    s = SDT_TRONG_CHU.sub("[SĐT]", s or "")
    s = EMAIL.sub("[email]", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s if len(s) <= toi_da else s[: toi_da - 1].rstrip() + "…"
