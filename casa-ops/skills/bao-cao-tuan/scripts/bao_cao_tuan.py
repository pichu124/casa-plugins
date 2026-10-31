#!/usr/bin/env python3
"""Báo cáo lead hằng tuần của Casa dei Bambini từ tab Leads của Sheet bot Messenger.

Đầu vào: file CSV (hoặc XLSX nếu có openpyxl) xuất từ tab Leads, hàng 1 là tiêu đề cột
(tiêu đề tiếng Việt như trên Sheet, hoặc khóa JSON của bot như `first_contact`).
Đầu ra: báo cáo Markdown (mặc định) hoặc JSON. Báo cáo không in SĐT đầy đủ, không in Customer ID.

Chỉ dùng thư viện chuẩn của Python. Mọi giờ là giờ Việt Nam, không múi giờ.

Ví dụ:
    python3 bao_cao_tuan.py leads.csv
    python3 bao_cao_tuan.py leads.csv --tuan 2026-10-28 --now "2026-11-02 08:00"
    python3 bao_cao_tuan.py leads.csv --format json
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import unicodedata
from collections import Counter, OrderedDict
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path

GIO_CHO_TOI_DA = 48  # Lead có SĐT quá chừng này giờ chưa liên hệ thì đưa vào danh sách.
HAU_TO_DA_LIEN_HE = "Đã liên hệ"

# Tiêu đề cột trên Sheet ↔ khóa JSON bot gửi (docs/google-sheet-leads.md của repo bot).
COT = OrderedDict(
    [
        ("time", "Thời gian"),
        ("event", "Sự kiện"),
        ("fb_name", "Tên Facebook"),
        ("customer_id", "Customer ID"),
        ("phone", "SĐT"),
        ("kid_name", "Tên bé"),
        ("kid_age", "Tuổi bé"),
        ("campus", "Cơ sở/hệ"),
        ("area", "Khu vực"),
        ("expected_start", "Dự kiến nhập học"),
        ("visit_request", "Hẹn tham quan/học thử"),
        ("ad_title", "Quảng cáo"),
        ("stage", "Giai đoạn"),
        ("last_message", "Tin cuối của phụ huynh"),
        ("first_contact", "Lần đầu liên hệ"),
        ("updated_at", "Cập nhật lúc"),
    ]
)

GIAI_DOAN = [
    "Mới",
    "Đã biết tuổi/cơ sở",
    "Đã có SĐT",
    "Muốn hẹn tham quan/học thử",
    "Đã hẹn tham quan/học thử",
    "Không quan tâm",
]
DA_HEN = "Đã hẹn tham quan/học thử"
MUON_HEN = "Muốn hẹn tham quan/học thử"
KHONG_QUAN_TAM = "Không quan tâm"

CO_SO = ["Chùa Láng", "Mỹ Đình", "Kim Mã", "Quốc tế – chưa rõ cơ sở", "Chưa rõ"]
NHOM_TUOI = [
    "15 tháng – <2 tuổi",
    "2 – <3 tuổi",
    "3 – <5 tuổi",
    "5 – 6 tuổi",
    "Ngoài độ tuổi",
    "Chưa rõ",
]


# ---------------------------------------------------------------- chuẩn hóa chữ


def bo_dau(s: str) -> str:
    """Chữ thường, bỏ dấu tiếng Việt, gộp khoảng trắng: để so khớp mềm."""
    s = unicodedata.normalize("NFD", s or "")
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "D").lower()
    return re.sub(r"\s+", " ", s).strip()


def doc_gio(s: str) -> datetime | None:
    s = (s or "").strip()
    if not s:
        return None
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%dT%H:%M:%S",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%Y-%m-%d",
        "%d/%m/%Y",
    ):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def chu_so(s: str) -> str:
    return re.sub(r"\D", "", s or "")


def co_sdt(s: str) -> bool:
    return len(chu_so(s)) >= 9


def che_sdt(s: str) -> str:
    so = chu_so(s)
    return f"…{so[-4:]}" if len(so) >= 4 else "…"


def ten_hien_thi(ten: str) -> str:
    ten = re.sub(r"\s+", " ", (ten or "")).strip()
    if not ten or bo_dau(ten) == "khach chua ro ten":
        return "Khách chưa rõ tên"
    return ten


def tach_giai_doan(gia_tri: str) -> tuple[str, bool]:
    """'Đã có SĐT · Đã liên hệ' → ('Đã có SĐT', True)."""
    phan = [p.strip() for p in re.split(r"[·•]", gia_tri or "") if p.strip()]
    da_lien_he = any(bo_dau(p) == bo_dau(HAU_TO_DA_LIEN_HE) for p in phan)
    chinh = [p for p in phan if bo_dau(p) != bo_dau(HAU_TO_DA_LIEN_HE)]
    ten = chinh[0] if chinh else ""
    for g in GIAI_DOAN:
        if bo_dau(g) == bo_dau(ten):
            return g, da_lien_he
    # Chữ cũ của bot ("Hẹn tham quan/học thử") là lịch AI hiểu, chưa chốt.
    if bo_dau(ten) == bo_dau("Hẹn tham quan/học thử"):
        return MUON_HEN, da_lien_he
    return (ten or "Chưa rõ"), da_lien_he


def chuan_hoa_co_so(s: str) -> str:
    t = bo_dau(s)
    if not t:
        return "Chưa rõ"
    nhom = {
        "Chùa Láng": ["chua lang", "hongkong", "hong kong", "huynh thuc khang", "hk"],
        "Mỹ Đình": ["my dinh", "flc", "pham hung"],
        "Kim Mã": ["kim ma", "van phuc", "song ngu", "ngoc ha"],
    }
    trung = [
        ten
        for ten, tu in nhom.items()
        if any(re.search(rf"(?<![a-z]){re.escape(x)}(?![a-z])", t) for x in tu)
    ]
    if len(trung) == 1:
        return trung[0]
    if len(trung) > 1:
        return "Chưa rõ"  # Phụ huynh nhắc nhiều cơ sở: không đoán.
    if re.search(r"quoc te|international", t):
        return "Quốc tế – chưa rõ cơ sở"
    return "Chưa rõ"


# ---------------------------------------------------------------- tuổi bé

SO_CHU = {"mot": 1, "hai": 2, "ba": 3, "bon": 4, "tu": 4, "nam": 5, "sau": 6, "bay": 7}


def _so(x: str) -> float:
    x = x.replace(",", ".")
    return float(SO_CHU[x]) if x in SO_CHU else float(x)


def khoang_thang(tuoi: str, ngay_tinh: datetime) -> tuple[float, float] | None:
    """Khoảng tuổi (tháng) [thấp, cao) mà chữ phụ huynh viết có thể ứng với. None nếu không hiểu."""
    t = bo_dau(tuoi)
    if not t:
        return None
    so = r"(\d+(?:[.,]\d+)?|mot|hai|ba|bon|tu|nam|sau|bay)"

    # Năm sinh: "sinh năm 2022", "sn 2022", "2022"; có tháng: "sinh tháng 3/2023", "03/2023".
    m = re.search(r"(?:thang\s*)?(\d{1,2})\s*[/-]\s*(20\d{2})", t)
    if m:
        thang, nam = int(m.group(1)), int(m.group(2))
        if 1 <= thang <= 12:
            tuoi_thang = (ngay_tinh.year - nam) * 12 + (ngay_tinh.month - thang)
            return (tuoi_thang, tuoi_thang + 1)
    m = re.search(r"(?<!\d)(20\d{2})(?!\d)", t)
    if m:
        nam = int(m.group(1))
        cao = (ngay_tinh.year - nam) * 12 + ngay_tinh.month  # sinh đầu năm
        thap = cao - 12  # sinh cuối năm
        return (max(thap, 0), cao)

    duoi = bool(re.search(r"\b(duoi|chua (du|toi|den))\b", t))

    # "3-4 tuổi", "3 đến 4 tuổi"
    m = re.search(rf"{so}\s*(?:-|–|den|toi)\s*{so}\s*tuoi", t)
    if m:
        a, b = _so(m.group(1)), _so(m.group(2))
        return (a * 12, (b + 1) * 12)

    # "x tuổi y tháng", "x tuổi rưỡi", "x tuổi"; "2t" là 2 tuổi ("18th" là tháng, xử lý dưới).
    t = re.sub(r"(\d)\s*t\b", r"\1 tuoi", t)
    m = re.search(rf"{so}\s*tuoi(?:\s*(ruoi)|\s*{so}\s*thang)?", t)
    if m:
        nam = _so(m.group(1))
        if m.group(2):
            return (nam * 12 + 6, nam * 12 + 7)
        if m.group(3):
            thang = _so(m.group(3))
            return (nam * 12 + thang, nam * 12 + thang + 1)
        if duoi:
            return (max(nam * 12 - 12, 0), nam * 12)
        if nam != int(nam):  # "1,5 tuổi"
            return (nam * 12, nam * 12 + 1)
        return (nam * 12, nam * 12 + 12)

    # "18 tháng", "dưới 18 tháng"
    m = re.search(rf"{so}\s*(?:thang|th\b)", t)
    if m:
        thang = _so(m.group(1))
        if duoi:
            return (max(thang - 6, 0), thang)
        return (thang, thang + 1)
    return None


def nhom_tuoi(tuoi: str, ngay_tinh: datetime) -> str:
    khoang = khoang_thang(tuoi, ngay_tinh)
    if khoang is None:
        return "Chưa rõ"
    giua = (khoang[0] + khoang[1]) / 2  # Lấy điểm giữa của khoảng có thể.
    # Trường nhận bé tới 6 tuổi: "6 tuổi" (72–83 tháng) vẫn thuộc nhóm 5 – 6 tuổi.
    if giua < 15 or giua >= 84:
        return "Ngoài độ tuổi"
    if giua < 24:
        return "15 tháng – <2 tuổi"
    if giua < 36:
        return "2 – <3 tuổi"
    if giua < 60:
        return "3 – <5 tuổi"
    return "5 – 6 tuổi"


# ---------------------------------------------------------------- đọc dữ liệu


@dataclass
class Lead:
    ten: str
    sdt: str
    su_kien: str
    giai_doan: str
    da_lien_he: bool
    co_so: str
    nhom_tuoi: str
    quang_cao: str
    lan_dau: datetime
    cap_nhat: datetime | None

    @property
    def co_sdt(self) -> bool:
        return co_sdt(self.sdt)

    def moc_co_sdt(self) -> datetime:
        """Lúc có SĐT. Sheet không có cột riêng: sự kiện mới nhất là SĐT mới/Đổi SĐT thì
        dùng 'Cập nhật lúc', không thì dùng 'Lần đầu liên hệ' (sớm hơn, nên tính dư giờ chờ)."""
        if self.cap_nhat and bo_dau(self.su_kien) in (bo_dau("SĐT mới"), bo_dau("Đổi SĐT")):
            return self.cap_nhat
        return self.lan_dau


@dataclass
class DuLieu:
    leads: list[Lead] = field(default_factory=list)
    hang_loi_ngay: int = 0
    trung_id: int = 0


def doc_hang(duong_dan: Path) -> list[dict[str, str]]:
    if duong_dan.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl  # type: ignore
        except ImportError:
            sys.exit("Cần openpyxl để đọc .xlsx (pip install openpyxl), hoặc xuất tab Leads ra .csv.")
        ws = openpyxl.load_workbook(duong_dan, read_only=True, data_only=True)
        sheet = ws["Leads"] if "Leads" in ws.sheetnames else ws.worksheets[0]
        hang = list(sheet.iter_rows(values_only=True))
        if not hang:
            return []
        tieu_de = ["" if c is None else str(c) for c in hang[0]]

        def o(c):
            if c is None:
                return ""
            if isinstance(c, datetime):
                return c.strftime("%Y-%m-%d %H:%M:%S")
            return str(c)

        return [dict(zip(tieu_de, (o(c) for c in r))) for r in hang[1:]]
    with open(duong_dan, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def lay(hang: dict[str, str], khoa: str) -> str:
    """Lấy ô theo tiêu đề tiếng Việt hoặc khóa JSON, không phân biệt hoa thường, dấu."""
    muc_tieu = {bo_dau(COT[khoa]), bo_dau(khoa)}
    for k, v in hang.items():
        if bo_dau(k or "") in muc_tieu:
            return (v or "").strip()
    return ""


def doc_du_lieu(hang: list[dict[str, str]], ngay_tinh: datetime) -> DuLieu:
    du_lieu = DuLieu()
    da_gap: set[str] = set()
    for h in hang:
        if not any((v or "").strip() for v in h.values()):
            continue
        cid = lay(h, "customer_id")
        if cid:
            if cid in da_gap:
                du_lieu.trung_id += 1
                continue
            da_gap.add(cid)
        lan_dau = doc_gio(lay(h, "first_contact"))
        if lan_dau is None:
            du_lieu.hang_loi_ngay += 1
            continue
        giai_doan, da_lien_he = tach_giai_doan(lay(h, "stage"))
        du_lieu.leads.append(
            Lead(
                ten=ten_hien_thi(lay(h, "fb_name")),
                sdt=lay(h, "phone"),
                su_kien=lay(h, "event"),
                giai_doan=giai_doan,
                da_lien_he=da_lien_he,
                co_so=chuan_hoa_co_so(lay(h, "campus")),
                nhom_tuoi=nhom_tuoi(lay(h, "kid_age"), ngay_tinh),
                quang_cao=re.sub(r"\s+", " ", lay(h, "ad_title")),
                lan_dau=lan_dau,
                cap_nhat=doc_gio(lay(h, "updated_at")),
            )
        )
    return du_lieu


# ---------------------------------------------------------------- tính chỉ số


def dau_tuan(d: datetime) -> datetime:
    d = d.replace(hour=0, minute=0, second=0, microsecond=0)
    return d - timedelta(days=d.weekday())


def ty_le(tu: int, mau: int) -> float | None:
    return None if mau == 0 else tu / mau


def chi_so_tuan(leads: list[Lead], bat_dau: datetime) -> dict:
    ket_thuc = bat_dau + timedelta(days=7)
    moi = [l for l in leads if bat_dau <= l.lan_dau < ket_thuc]
    co_so_dt = [l for l in moi if l.co_sdt]
    da_hen = [l for l in co_so_dt if l.giai_doan == DA_HEN]
    muon_hen = [l for l in co_so_dt if l.giai_doan == MUON_HEN]
    tu_qc = [l for l in moi if l.quang_cao]

    giai_doan = Counter(l.giai_doan for l in moi)
    theo_quang_cao = {}
    for ten, so in Counter(l.quang_cao for l in tu_qc).most_common():
        co = sum(1 for l in tu_qc if l.quang_cao == ten and l.co_sdt)
        hen = sum(1 for l in tu_qc if l.quang_cao == ten and l.co_sdt and l.giai_doan == DA_HEN)
        theo_quang_cao[ten] = {"lead": so, "co_sdt": co, "ty_le_co_sdt": ty_le(co, so), "da_hen": hen}

    return {
        "bat_dau": bat_dau.strftime("%Y-%m-%d"),
        "ket_thuc": (ket_thuc - timedelta(days=1)).strftime("%Y-%m-%d"),
        "lead_moi": len(moi),
        "theo_giai_doan": {g: giai_doan.get(g, 0) for g in GIAI_DOAN}
        | {g: n for g, n in giai_doan.items() if g not in GIAI_DOAN},
        "da_lien_he": sum(1 for l in moi if l.da_lien_he),
        "theo_co_so": {c: sum(1 for l in moi if l.co_so == c) for c in CO_SO},
        "theo_tuoi": {n: sum(1 for l in moi if l.nhom_tuoi == n) for n in NHOM_TUOI},
        "co_sdt": len(co_so_dt),
        "da_hen": len(da_hen),
        "muon_hen_chua_chot": len(muon_hen),
        "ty_le_sdt_hen": ty_le(len(da_hen), len(co_so_dt)),
        "lead_tu_quang_cao": len(tu_qc),
        "theo_quang_cao": theo_quang_cao,
    }


def cho_lien_he(leads: list[Lead], bay_gio: datetime) -> list[dict]:
    ds = []
    for l in leads:
        if not l.co_sdt or l.da_lien_he or l.giai_doan in (DA_HEN, KHONG_QUAN_TAM):
            continue
        gio = (bay_gio - l.moc_co_sdt()).total_seconds() / 3600
        if gio > GIO_CHO_TOI_DA:
            ds.append(
                {
                    "ten": l.ten,
                    "sdt_4_so": che_sdt(l.sdt),
                    "co_so": l.co_so,
                    "nhom_tuoi": l.nhom_tuoi,
                    "giai_doan": l.giai_doan,
                    "gio_cho": int(gio),
                }
            )
    return sorted(ds, key=lambda x: -x["gio_cho"])


def chenh(nay: float | None, truoc: float | None) -> dict:
    if nay is None or truoc is None:
        return {"nay": nay, "truoc": truoc, "chenh": None, "phan_tram": None}
    return {
        "nay": nay,
        "truoc": truoc,
        "chenh": nay - truoc,
        "phan_tram": None if truoc == 0 else (nay - truoc) / truoc,
    }


def lap_bao_cao(du_lieu: DuLieu, bat_dau: datetime, bay_gio: datetime) -> dict:
    tuan = chi_so_tuan(du_lieu.leads, bat_dau)
    truoc = chi_so_tuan(du_lieu.leads, bat_dau - timedelta(days=7))
    return {
        "chay_luc": bay_gio.strftime("%Y-%m-%d %H:%M"),
        "tuan": tuan,
        "tuan_truoc": truoc,
        "so_sanh": {
            "lead_moi": chenh(tuan["lead_moi"], truoc["lead_moi"]),
            "ty_le_sdt_hen": chenh(tuan["ty_le_sdt_hen"], truoc["ty_le_sdt_hen"]),
            "lead_tu_quang_cao": chenh(tuan["lead_tu_quang_cao"], truoc["lead_tu_quang_cao"]),
        },
        "cho_lien_he_qua_48h": cho_lien_he(du_lieu.leads, bay_gio),
        "ghi_chu_du_lieu": {
            "tong_hang_doc_duoc": len(du_lieu.leads),
            "hang_bo_qua_loi_ngay": du_lieu.hang_loi_ngay,
            "hang_trung_customer_id": du_lieu.trung_id,
            "lead_moi_chua_ro_co_so": tuan["theo_co_so"]["Chưa rõ"],
            "lead_moi_chua_ro_tuoi": tuan["theo_tuoi"]["Chưa rõ"],
        },
    }


# ---------------------------------------------------------------- trình bày


def pt(x: float | None) -> str:
    if x is None:
        return "–"
    s = f"{x * 100:.1f}".replace(".", ",")
    return (s[:-2] if s.endswith(",0") else s) + "%"


def ngay(s: str) -> str:
    return datetime.strptime(s, "%Y-%m-%d").strftime("%d/%m/%Y")


def dau(x: float) -> str:
    return f"+{x:g}" if x > 0 else f"{x:g}"


def mo_ta_chenh(c: dict, la_ty_le: bool = False) -> str:
    if c["chenh"] is None:
        return "chưa so được"
    if la_ty_le:
        diem = c["chenh"] * 100
        return f"{dau(round(diem, 1)).replace('.', ',')} điểm %"
    s = dau(c["chenh"])
    return s if c["phan_tram"] is None else f"{s} ({dau(round(c['phan_tram'] * 100, 1)).replace('.', ',')}%)"


def bang(tieu_de: list[str], hang: list[list]) -> str:
    dong = ["| " + " | ".join(tieu_de) + " |", "|" + "|".join("---" for _ in tieu_de) + "|"]
    dong += ["| " + " | ".join(str(o) for o in h) + " |" for h in hang]
    return "\n".join(dong)


def markdown(bc: dict) -> str:
    t, tr, ss = bc["tuan"], bc["tuan_truoc"], bc["so_sanh"]
    out = [
        f"# Báo cáo lead tuần {ngay(t['bat_dau'])} – {ngay(t['ket_thuc'])}",
        "",
        f"_Lập lúc {datetime.strptime(bc['chay_luc'], '%Y-%m-%d %H:%M').strftime('%H:%M %d/%m/%Y')}, "
        "từ tab Leads của Sheet bot Messenger._",
        "",
        "## Tóm tắt",
        "",
        f"- **{t['lead_moi']} lead mới**, tuần trước {tr['lead_moi']}, chênh {mo_ta_chenh(ss['lead_moi'])}.",
        f"- **Có SĐT → đã hẹn: {pt(t['ty_le_sdt_hen'])}** ({t['da_hen']}/{t['co_sdt']} lead có SĐT đã chốt lịch), "
        f"tuần trước {pt(tr['ty_le_sdt_hen'])}, chênh {mo_ta_chenh(ss['ty_le_sdt_hen'], True)}. "
        f"Thêm {t['muon_hen_chua_chot']} lead muốn hẹn nhưng chưa chốt.",
        f"- **{t['lead_tu_quang_cao']} lead từ quảng cáo**, tuần trước {tr['lead_tu_quang_cao']}, "
        f"chênh {mo_ta_chenh(ss['lead_tu_quang_cao'])}.",
        f"- **{len(bc['cho_lien_he_qua_48h'])} lead có SĐT quá {GIO_CHO_TOI_DA} giờ chưa liên hệ** (danh sách cuối báo cáo).",
        "",
        "## Theo giai đoạn (lead mới trong tuần)",
        "",
        bang(["Giai đoạn", "Số lead"], [[g, n] for g, n in t["theo_giai_doan"].items()]),
        "",
        f"Trong đó {t['da_lien_he']} lead đã được tư vấn viên đánh dấu đã liên hệ.",
        "",
        "## Theo cơ sở",
        "",
        bang(["Cơ sở", "Số lead"], [[c, n] for c, n in t["theo_co_so"].items()]),
        "",
        "## Theo độ tuổi",
        "",
        bang(["Nhóm tuổi", "Số lead"], [[c, n] for c, n in t["theo_tuoi"].items()]),
        "",
        "## Lead từ quảng cáo",
        "",
    ]
    if t["theo_quang_cao"]:
        out.append(
            bang(
                ["Quảng cáo", "Lead", "Có SĐT", "Tỷ lệ có SĐT", "Đã hẹn"],
                [
                    [ten, q["lead"], q["co_sdt"], pt(q["ty_le_co_sdt"]), q["da_hen"]]
                    for ten, q in t["theo_quang_cao"].items()
                ],
            )
        )
    else:
        out.append("Tuần này không có lead nào từ quảng cáo.")
    out += [
        "",
        "## So với tuần trước",
        "",
        bang(
            ["Chỉ số", "Tuần này", "Tuần trước", "Chênh lệch"],
            [
                ["Lead mới", t["lead_moi"], tr["lead_moi"], mo_ta_chenh(ss["lead_moi"])],
                ["Có SĐT → đã hẹn", pt(t["ty_le_sdt_hen"]), pt(tr["ty_le_sdt_hen"]), mo_ta_chenh(ss["ty_le_sdt_hen"], True)],
                ["Lead từ quảng cáo", t["lead_tu_quang_cao"], tr["lead_tu_quang_cao"], mo_ta_chenh(ss["lead_tu_quang_cao"])],
            ],
        ),
        "",
        f"## Lead có SĐT quá {GIO_CHO_TOI_DA} giờ chưa liên hệ",
        "",
    ]
    ds = bc["cho_lien_he_qua_48h"]
    if ds:
        out.append(
            bang(
                ["Phụ huynh", "Cơ sở", "Tuổi bé", "Giai đoạn", "Đã chờ"],
                [
                    [f"{x['ten']} ({x['sdt_4_so']})", x["co_so"], x["nhom_tuoi"], x["giai_doan"], f"{x['gio_cho']} giờ"]
                    for x in ds
                ],
            )
        )
        out += [
            "",
            "Tính mọi lead, không riêng tuần này. Giờ chờ tính từ lúc có SĐT; khi Sheet không ghi được mốc đó "
            "thì tính từ lần đầu liên hệ, nên có thể dư.",
        ]
    else:
        out.append("Không có lead nào.")
    g = bc["ghi_chu_du_lieu"]
    ghi_chu = []
    if g["hang_bo_qua_loi_ngay"]:
        ghi_chu.append(f"{g['hang_bo_qua_loi_ngay']} hàng bị bỏ qua vì không đọc được ngày lần đầu liên hệ.")
    if g["hang_trung_customer_id"]:
        ghi_chu.append(f"{g['hang_trung_customer_id']} hàng trùng Customer ID, chỉ tính hàng đầu tiên.")
    if g["lead_moi_chua_ro_co_so"]:
        ghi_chu.append(f"{g['lead_moi_chua_ro_co_so']} lead mới chưa rõ cơ sở.")
    if g["lead_moi_chua_ro_tuoi"]:
        ghi_chu.append(f"{g['lead_moi_chua_ro_tuoi']} lead mới chưa rõ tuổi bé.")
    if ghi_chu:
        out += ["", "## Ghi chú về dữ liệu", ""] + [f"- {x}" for x in ghi_chu]
    return "\n".join(out) + "\n"


# ---------------------------------------------------------------- chạy


def gio_viet_nam() -> datetime:
    try:
        from zoneinfo import ZoneInfo

        return datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).replace(tzinfo=None)
    except Exception:
        return datetime.utcnow() + timedelta(hours=7)


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Báo cáo lead tuần của Casa dei Bambini.")
    p.add_argument("file", type=Path, help="CSV hoặc XLSX xuất từ tab Leads")
    p.add_argument("--tuan", help="Một ngày bất kỳ trong tuần cần báo cáo (YYYY-MM-DD). Mặc định: tuần trước.")
    p.add_argument("--now", help="Giờ chạy báo cáo, giờ Việt Nam (YYYY-MM-DD HH:MM). Mặc định: bây giờ.")
    p.add_argument("--format", choices=["markdown", "json"], default="markdown")
    a = p.parse_args(argv)

    bay_gio = doc_gio(a.now) if a.now else gio_viet_nam()
    if bay_gio is None:
        p.error("--now không đúng dạng YYYY-MM-DD HH:MM")
    if a.tuan:
        ngay_tuan = doc_gio(a.tuan)
        if ngay_tuan is None:
            p.error("--tuan không đúng dạng YYYY-MM-DD")
        bat_dau = dau_tuan(ngay_tuan)
    else:
        bat_dau = dau_tuan(bay_gio) - timedelta(days=7)

    du_lieu = doc_du_lieu(doc_hang(a.file), bay_gio)
    bc = lap_bao_cao(du_lieu, bat_dau, bay_gio)
    if a.format == "json":
        print(json.dumps(bc, ensure_ascii=False, indent=2))
    else:
        print(markdown(bc), end="")
    return 0


if __name__ == "__main__":
    sys.exit(main())
